#!/usr/bin/env python3
"""Build the anonymous Nongnet idea proposal from the retained official form."""
from copy import deepcopy
from pathlib import Path
from datetime import datetime, timezone, timedelta
import argparse
import hashlib
import json
import os
import re
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / 'reference' / 'nongnet-idea-template.docx'
PROJECT = ROOT.parent
EXAMPLE_FORECAST = {
    'artifact': 'outputs/prospective-e106/2026-10-11/validated-runtime/forecast.json',
    'artifactSha256': 'bd2e8566d15aed3e8a4765cd43d7def35a9b7ee04a573d18ff9a426f6364cbdc',
    'canonicalSha256': 'dbc79c9624bb5cacc4de8c3477c3b253f0bec0ebe0c3e81a2bba816102f9873f',
    'issueDateKst': '2026-10-11',
    'issueAtKst': '2026-10-11T15:19:46+09:00',
    'dailyTrainingRunId': '20261011T120011-0014dce9',
    'dailyTrainingSnapshotRecordSha256': '2d5a668e8d402f0ba4ad907e3a54756b40962b926c076a9b7f108ad184526ea2',
    'runStatusFileSha256': 'ce9cbf87db731a52e1233b07ef1aede6f750747477ac0564a7ca7fac6c718683',
    'sourceCutoffDate': '2026-10-10',
    'pricesFileSha256': 'd5fc7919991cd8ac5ff6574de097d8697a60cdf2aa4d8ae22c23c95d6160930b',
    'provenanceFileSha256': '6fde95eb6fdd9c76febcd9e05d8f6814cf7273c731e5cebb0d6b023dcdc3e667',
    'provenanceCanonicalSha256': 'cdb9bb5e1471c6358cb9db45abb179756be33493d1c77dabc54d10f63cf721c1',
    'selectionFileSha256': '59fa52c3f7c9664d52ad8ce23987eb4d6c70af0fca9907c25596a84eb77ba26b',
}

parser = argparse.ArgumentParser(description='Build the anonymous Nongnet idea proposal.')
parser.add_argument('--receipt', help='official receipt number (ASCII letters, digits, underscore, or hyphen)')
args = parser.parse_args()
if args.receipt is not None and not re.fullmatch(r'[A-Za-z0-9_-]{1,40}', args.receipt):
    parser.error('receipt must be 1–40 ASCII letters, digits, underscores, or hyphens')
OUTPUT = ROOT / (f'아이디어예측활용안_{args.receipt}.docx' if args.receipt is not None else 'idea-proposal-v1.docx')

NAVY = '315A98'
PALE = 'EDF2F8'
GRID = 'D9D9D9'
BLACK = '000000'
FONT = 'Nanum Gothic'


def set_font(run, size=10.5, bold=None, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn('w:' + key), FONT)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcpr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_borders(cell, color=GRID, size='5'):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcpr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tcpr.append(margins)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            margins.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    trpr.append(tbl_header)


def cell_text(cell, text, *, header=False, size=9.5, align=None, bold=False):
    cell.text = ''
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=(True if header else bold), color=('FFFFFF' if header else BLACK))
    set_cell_shading(cell, NAVY if header else 'FFFFFF')
    set_cell_borders(cell)
    set_cell_margins(cell)


def style_table(table, header=True, size=9.2, widths=None):
    table.autofit = False
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    for ri, row in enumerate(table.rows):
        if ri == 0 and header:
            set_repeat_table_header(row)
        for ci, cell in enumerate(row.cells):
            txt = cell.text
            cell_text(cell, txt, header=(ri == 0 and header), size=size,
                      align=WD_ALIGN_PARAGRAPH.CENTER if ci == 0 and len(row.cells) <= 4 else None)
            if ri > 0 and ri % 2 == 0:
                set_cell_shading(cell, PALE)


def add_text(doc, text='', *, size=10.5, bold=False, before=0, after=6, keep=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.13
    p.paragraph_format.keep_together = keep
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold)
    return p


def add_label(doc, text):
    p = add_text(doc, text, size=11, bold=True, before=6, after=4, keep=True)
    return p


def add_table(doc, rows, widths, *, size=9.2, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell_text(table.cell(ri, ci), value, header=(header and ri == 0), size=size)
            if ri > 0 and ri % 2 == 0:
                set_cell_shading(table.cell(ri, ci), PALE)
    style_table(table, header=header, size=size, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def section_banner(doc, original):
    clone = deepcopy(original._tbl)
    body = doc._body._body
    body.insert(len(body) - 1, clone)
    # Keep adjacent full-width section banners from merging into one table.
    spacer = OxmlElement('w:p')
    body.insert(len(body) - 1, spacer)


def set_all_existing_fonts(doc):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_font(run, size=(run.font.size.pt if run.font.size else 10.5), bold=run.bold)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_font(run, size=(run.font.size.pt if run.font.size else 10), bold=run.bold)


def load_verified_radish_example():
    """Read the pinned prospective E106 issue and verify its immutable lineage."""
    index = EXAMPLE_FORECAST
    artifact_path = (PROJECT / index['artifact']).resolve()
    if PROJECT.resolve() not in artifact_path.parents:
        raise SystemExit('Forecast artifact path escapes the project directory.')
    raw = artifact_path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != index['artifactSha256']:
        raise SystemExit('Pinned example artifact hash mismatch; stop before using example values.')
    issue = json.loads(raw)
    unsigned = dict(issue)
    embedded_canonical = unsigned.pop('forecastSha256', None)
    canonical_raw = (json.dumps(unsigned, ensure_ascii=False, sort_keys=True,
                                separators=(',', ':'), allow_nan=False) + '\n').encode('utf-8')
    canonical_sha = hashlib.sha256(canonical_raw).hexdigest()
    if canonical_sha != index['canonicalSha256'] or embedded_canonical != canonical_sha:
        raise SystemExit('Pinned example canonical forecast hash mismatch.')
    sources = issue.get('sources')
    if not isinstance(sources, dict):
        raise SystemExit('Forecast source lineage is missing.')
    if (issue.get('issueDateKst') != index['issueDateKst'] or
            issue.get('issueAtKst') != index['issueAtKst'] or
            sources.get('dailyTrainingRunId') != index['dailyTrainingRunId'] or
            sources.get('dailyTrainingSnapshotRecordSha256') != index['dailyTrainingSnapshotRecordSha256'] or
            sources.get('dailyTrainingFinishedAtKst') is None or
            sources.get('sourceCutoff') != index['sourceCutoffDate'] or
            sources.get('pricesFileSha256') != index['pricesFileSha256'] or
            sources.get('provenanceFileSha256') != index['provenanceFileSha256'] or
            sources.get('provenanceCanonicalSha256') != index['provenanceCanonicalSha256'] or
            sources.get('selectionFileSha256') != index['selectionFileSha256']):
        raise SystemExit('Forecast issue/index provenance does not agree.')
    if (issue.get('schemaVersion') != 1 or issue.get('phase') != 'prospective_forecast_only' or
            issue.get('targetValuesAttached') is not False or
            issue.get('methods') != ['causal50', 'E', 'E106']):
        raise SystemExit('Pinned issue does not match the answer-free causal50/E/E106 protocol.')
    def project_file(relative, label):
        path = (PROJECT / relative).resolve()
        if PROJECT.resolve() not in path.parents:
            raise SystemExit(f'{label} path escapes the project directory.')
        return path

    prices_path = project_file(sources.get('pricesPath'), 'Price snapshot')
    prices_raw = prices_path.read_bytes()
    if hashlib.sha256(prices_raw).hexdigest() != index['pricesFileSha256']:
        raise SystemExit('Pinned price snapshot file hash mismatch.')
    prices = json.loads(prices_raw)
    provenance_path = project_file(sources.get('provenancePath'), 'Selected provenance')
    provenance_raw = provenance_path.read_bytes()
    provenance = json.loads(provenance_raw)
    provenance_canonical = (json.dumps(provenance, ensure_ascii=False, sort_keys=True,
                                      separators=(',', ':'), allow_nan=False) + '\n').encode('utf-8')
    if (hashlib.sha256(provenance_raw).hexdigest() != index['provenanceFileSha256'] or
            hashlib.sha256(provenance_canonical).hexdigest() != index['provenanceCanonicalSha256']):
        raise SystemExit('Pinned selected-provenance file/canonical hash mismatch.')
    if (prices_path.parent != provenance_path.parent or prices_path.name != 'prices.json' or
            provenance_path.name != 'selected-provenance.json' or
            prices_path.parent.name != index['dailyTrainingRunId']):
        raise SystemExit('Price and selected-provenance files are not from the pinned daily-training run.')
    selection_path = project_file(
        'outputs/experiments/trend-response-correction-20261010/selection.json', 'Trend selection')
    if hashlib.sha256(selection_path.read_bytes()).hexdigest() != index['selectionFileSha256']:
        raise SystemExit('Pinned E selection file hash mismatch.')
    run_status_path = prices_path.parent / 'run-status.json'
    run_status_raw = run_status_path.read_bytes()
    if hashlib.sha256(run_status_raw).hexdigest() != index['runStatusFileSha256']:
        raise SystemExit('Pinned immutable daily-training run-status hash mismatch.')
    run_status = json.loads(run_status_raw)
    if (run_status.get('status') != 'complete' or run_status.get('runId') != index['dailyTrainingRunId'] or
            run_status.get('snapshotSha256') != index['dailyTrainingSnapshotRecordSha256'] or
            run_status.get('cutoff') != index['sourceCutoffDate']):
        raise SystemExit('Pinned source run has no matching complete run-status record.')
    if (run_status.get('finishedAtKst') != sources.get('dailyTrainingFinishedAtKst') or
            datetime.fromisoformat(run_status['finishedAtKst']) >
            datetime.fromisoformat(index['issueAtKst'])):
        raise SystemExit('Daily-training completion is missing, inconsistent, or after the E106 issue time.')
    if (issue.get('forecastSha256') != canonical_sha or
            sources.get('dailyTrainingCutoff') != index['sourceCutoffDate'] or
            sources.get('noTargetOutcomesRead') is not True):
        raise SystemExit('Daily-training lineage or answer-free marker is invalid.')

    issue_day = datetime.fromisoformat(index['issueDateKst']).date()
    cutoff_day = datetime.fromisoformat(index['sourceCutoffDate']).date()
    if cutoff_day >= issue_day or prices.get('meta', {}).get('evaluationEnd') != cutoff_day.isoformat():
        raise SystemExit('Pinned source cutoff does not agree with the price snapshot.')
    products = issue.get('products')
    if set(products or []) != {'cabbage', 'radish', 'onion', 'garlic'}:
        raise SystemExit('Forecast issue does not have the four-crop E106 domain.')
    forecasts = issue.get('forecasts')
    if not isinstance(forecasts, list) or len(forecasts) != 40:
        raise SystemExit('Forecast issue must contain exactly 40 forecast records.')
    expected_grid = {(product, horizon) for product in products for horizon in range(1, 11)}
    if any(not isinstance(row, dict) for row in forecasts):
        raise SystemExit('Forecast issue contains a malformed forecast record.')
    actual_grid = {(row.get('product'), row.get('horizon')) for row in forecasts}
    if len(actual_grid) != 40 or actual_grid != expected_grid:
        raise SystemExit('Forecast issue does not contain the exact four-crop, ten-horizon grid.')
    radish_rows = [x for x in forecasts or [] if x.get('product') == 'radish']
    if len(radish_rows) != 10:
        raise SystemExit('Radish issue does not contain exactly ten forecast horizons.')
    predictions = sorted([{'horizonDays': row['horizon'], 'target': row['target'],
                           'predictedPrice': row['E106']['prediction'],
                           'sourceDate': row.get('p0Date')}
                          for row in radish_rows if row.get('E106', {}).get('status') == 'issued'],
                         key=lambda x: x['horizonDays'])
    if [x['horizonDays'] for x in predictions] != list(range(1, 11)):
        raise SystemExit('Radish E106 does not issue all ten forecast horizons.')
    expected_targets = [(issue_day + timedelta(days=h)).isoformat() for h in range(1, 11)]
    if [x['target'] for x in predictions] != expected_targets:
        raise SystemExit('Radish target dates differ from the pinned E106 issue.')
    availability = issue.get('sourceAvailability', {}).get('radish')
    if not isinstance(availability, dict):
        raise SystemExit('Radish source availability audit is missing.')
    p0_dates = {row.get('p0Date') for row in radish_rows}
    p0_values = {row.get('p0') for row in radish_rows}
    if (len(p0_dates) != 1 or len(p0_values) != 1 or None in p0_dates or None in p0_values or
            availability.get('lastObservedPriceDate') != next(iter(p0_dates))):
        raise SystemExit('Radish E106 forecasts do not share the verified latest eligible source price.')
    product_prices = next((p for p in prices.get('products', []) if p.get('id') == 'radish'), None)
    actual = next((row for row in (product_prices or {}).get('observations', [])
                   if row.get('date') == next(iter(p0_dates))), None)
    if not actual or float(actual.get('actual', 0)) != float(next(iter(p0_values))):
        raise SystemExit('Radish forecast p0 does not match the pinned price snapshot observation.')
    source_states = [x for x in availability.get('sourceStates90d', [])
                     if x.get('date') == cutoff_day.isoformat()]
    if (len(source_states) != 1 or source_states[0].get('state') != 'complete_rows' or
            set(source_states[0].get('selectedSources') or {}) != {'110001', '110008'} or
            any(source_states[0]['selectedSources'][m]['rowCount'] <= 0 for m in ('110001', '110008'))):
        raise SystemExit('Radish cutoff provenance is not complete and nonempty for both POC markets.')
    provenance_cutoff = [row for row in provenance if row.get('product') == 'radish' and
                         row.get('date') == cutoff_day.isoformat()]
    if ({row.get('market') for row in provenance_cutoff} != {'110001', '110008'} or
            any(row.get('rowCount', 0) <= 0 for row in provenance_cutoff)):
        raise SystemExit('Pinned selected provenance does not contain both nonempty radish cutoff scopes.')
    selected_sources = source_states[0]['selectedSources']
    for market in ('110001', '110008'):
        row = next(x for x in provenance_cutoff if x['market'] == market)
        selected = selected_sources[market]
        if any(row.get(field) != selected.get(target) for field, target in
               (('rowCount', 'rowCount'), ('jobId', 'jobId'), ('epochId', 'epochId'),
                ('responseAt', 'responseAt'), ('selectedPrevious', 'selectedPrevious'))):
            raise SystemExit('Radish source-state audit does not match selected provenance bytes.')
    product = {'productId': 'radish', 'productName': '무', 'status': 'available', 'method': 'E106',
               'lastObservedPrice': float(next(iter(p0_values))),
               'lastObservedPriceDate': next(iter(p0_dates))}
    issue['sourceCutoffDate'] = index['sourceCutoffDate']
    return index, issue, product, predictions, None


def scrub_package_metadata(docx_path, expected_pages=3):
    """Remove inherited external links and empty note parts; normalize metadata."""
    ns = {
        'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
        'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
        'cp': 'http://schemas.openxmlformats.org/package/2006/metadata/core-properties',
        'dc': 'http://purl.org/dc/elements/1.1/',
        'dcterms': 'http://purl.org/dc/terms/',
        'xsi': 'http://www.w3.org/2001/XMLSchema-instance',
        'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties',
    }
    ET.register_namespace('cp', ns['cp'])
    ET.register_namespace('dc', ns['dc'])
    ET.register_namespace('dcterms', ns['dcterms'])
    ET.register_namespace('xsi', ns['xsi'])
    ET.register_namespace('vt', 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes')
    with ZipFile(docx_path, 'r') as src:
        parts = {name: src.read(name) for name in src.namelist()}

    rel_root = ET.fromstring(parts['word/_rels/document.xml.rels'])
    remove_rel_ids = set()
    for rel in list(rel_root):
        rel_type = rel.attrib.get('Type', '')
        target = rel.attrib.get('Target', '')
        if rel.attrib.get('TargetMode') == 'External' or rel_type.endswith(('/footnotes', '/endnotes')):
            remove_rel_ids.add(rel.attrib.get('Id'))
            rel_root.remove(rel)
    ET.register_namespace('', ns['rel'])
    parts['word/_rels/document.xml.rels'] = ET.tostring(rel_root, encoding='utf-8', xml_declaration=True)
    for name in ('word/footnotes.xml', 'word/endnotes.xml'):
        parts.pop(name, None)

    ct_root = ET.fromstring(parts['[Content_Types].xml'])
    for override in list(ct_root):
        if override.attrib.get('PartName') in ('/word/footnotes.xml', '/word/endnotes.xml'):
            ct_root.remove(override)
    ET.register_namespace('', ns['ct'])
    parts['[Content_Types].xml'] = ET.tostring(ct_root, encoding='utf-8', xml_declaration=True)

    core = ET.fromstring(parts['docProps/core.xml'])
    def core_child(tag):
        return core.find(tag)
    for tag, value in ((f"{{{ns['dc']}}}creator", ''), (f"{{{ns['cp']}}}lastModifiedBy", ''),
                       (f"{{{ns['cp']}}}keywords", ''), (f"{{{ns['dc']}}}description", '')):
        node = core_child(tag)
        if node is not None:
            node.text = value
    printed = core_child(f"{{{ns['cp']}}}lastPrinted")
    if printed is not None:
        core.remove(printed)
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace('+00:00', 'Z')
    for tag in (f"{{{ns['dcterms']}}}created", f"{{{ns['dcterms']}}}modified"):
        node = core_child(tag)
        if node is not None:
            node.text = now
    parts['docProps/core.xml'] = ET.tostring(core, encoding='utf-8', xml_declaration=True)

    if 'docProps/app.xml' in parts:
        app = ET.fromstring(parts['docProps/app.xml'])
        for child in list(app):
            local = child.tag.rsplit('}', 1)[-1]
            if local in {'TotalTime', 'Pages', 'Words', 'Characters', 'Lines', 'Paragraphs', 'HeadingPairs', 'TitlesOfParts'}:
                app.remove(child)
            elif local in {'Company', 'Manager'}:
                child.text = ''
        parts['docProps/app.xml'] = ET.tostring(app, encoding='utf-8', xml_declaration=True)

    temp_path = docx_path.with_name(docx_path.stem + '.scrubbed.docx')
    with ZipFile(temp_path, 'w', compression=ZIP_DEFLATED) as dst:
        for name, data in parts.items():
            dst.writestr(name, data)
    os.replace(temp_path, docx_path)


example_index, example_issue, radish, radish_predictions, thin_trade = load_verified_radish_example()

template_hash = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
expected_hash = 'd3f49cb015c186a0901dccf0496545c815b2b3fb270ea75dfe77780a81b01dec'
if template_hash != expected_hash:
    raise SystemExit('Retained template hash does not match artifact.md; stop and redistill.')

doc = Document(TEMPLATE)
source_banners = {2: deepcopy(doc.tables[2]), 5: deepcopy(doc.tables[5]), 9: deepcopy(doc.tables[9])}

# Remove source content while retaining the section properties and package styles.
body = doc._body._body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.top_margin = Inches(0.69)
section.bottom_margin = Inches(0.62)
for footer_paragraph in section.footer.paragraphs:
    footer_paragraph.clear()

doc.core_properties.title = '농산물 단기 가격 전망과 예측 근거 공개 활용안'
doc.core_properties.subject = '아이디어 예측 활용안'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.keywords = ''
doc.core_properties.comments = ''

p = doc.add_paragraph(style='Title')
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(5)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run('농산물 단기 가격 전망과 예측 근거 공개 활용안')
set_font(r, 16, bold=True)
receipt_label = args.receipt if args.receipt is not None else '[미발급]'
add_text(doc, f'접수번호 {receipt_label}    실증: 배추 · 무 · 양파 · 깐마늘    건고추: 확장 계획', size=10.5, after=12)

section_banner(doc, source_banners[2])
add_label(doc, '1-1  예측아이디어')
add_text(doc, '품목별 현재 가격과 향후 10일 전망을 자료 최신일·과거 오차와 함께 보여주는 대시보드를 제안합니다. 배추·무·양파·깐마늘의 가락·강서 경매가격을 대상으로, 가격 이력에서 미래 가격을 직접 예측하는 규제 회귀와 달력형 시계열 모형을 절반씩 결합했습니다. 배추의 추세 반응 보정과 깐마늘의 확장 가격 특징을 적용한 연구 후보 E106도 보존해 같은 정보 조건에서 비교합니다.', after=6)
add_text(doc, '2026년 7월11일~9월27일 후향 공통표본에서 네 품목·10거리 동일가중 MAPE는 기본 결합 14.7121%, 배추 보정 E 14.6500%, E106 14.6126%였습니다. E106의 MAE는 281.82원/kg으로 E의 285.74원/kg보다 낮았고, 발행 범위 3,160키와 평가 공통쌍 2,470개를 유지했습니다. 반복 열람한 자료의 작은 평균 개선이며, 통계적 우월성은 미확정입니다. 기본 결합 대비 MAPE 차이 −0.0995%p의 날짜 블록 95% 구간도 [−0.2556, +0.0599]%p로 0을 포함합니다.', after=6)
rhs_update_text = 'H/S는 V 가격 기준 미달로 중단해 T를 실행하지 않았습니다. RHS는 무의 V/T 결합 MAPE·MAE 기준을 통과했습니다. C−B 무 결합 MAPE 개선은 V 0.00276%p, T 0.01648%p였고, 4품목 T도 0.00412%p로 작았습니다. T 평균 블록구간은 개선 방향이나 거리별 검정은 미확정이며, 거리 순서의 추가 이점과 미래 성능은 미확정입니다. 구간별 월·거리 적합 30개의 정규화 비중 평균에서 C 계수 변화 제곱량 약98%가 작은 고유값 하위 10개 방향에 모였습니다. 결합 가격 이동은 작고 가격 이동량 대비 MAPE 순개선 비율은 V 2.32%, T 9.35%로 낮아 이득·손해가 상쇄됐습니다. 분산 감소의 인과 증거는 아닙니다. RHS는 후향 연구 대조군으로 보존하고 운영 E106은 유지합니다.'
rhs_update_p = add_text(doc, rhs_update_text, size=9.5, after=6)
if rhs_update_p.text != rhs_update_text:
    raise RuntimeError('RHS proposal paragraph was not added as expected.')
add_text(doc, '현재 실증은 두 시장·네 품목·10일 일별 전망입니다. 공식 전국가격·다섯 품목·2순/3순 성능과 구분하고, 공식 자료의 품종·단위·집계 규칙이 확인되면 별도로 연결합니다. 실제 미래 정확도는 같은 시각·스냅숏으로 발행한 기본 결합과 E·E106에 이후 정답을 연결해 확인합니다. 과거 원천의 최초 공개판을 입증하지 못하는 한계가 있어, 새 발행부터 입력 버전과 자료 상태를 고정 보존합니다.', after=9)

add_label(doc, '1-2  예측아이디어 도식')
add_text(doc, '가격 전망에 사용한 정보와 검증 범위를 함께 공개합니다.', size=9.5, after=4)
add_table(doc, [
    ['① 원천 자료 확인', '→', '② 가격이력 학습', '→', '③ 10일 전망 발행', '→', '④ 정답 비교·설명'],
    ['품목·시장·단위·수정 버전과 원천 처리 상태 확인', '→', '원점 이전 정보로 직접 회귀와 시계열 모형을 학습', '→', '결합·품목별 보정과 기준 전망을 같은 시점에 고정', '→', '과거 오차·자료 최신일과 후보별 미래 평가를 구분'],
], [1.06, 0.22, 1.40, 0.22, 1.45, 0.22, 1.40], size=8.2)
add_text(doc, '개선 후보가 실패해도 이전 개선안을 유지합니다. 새로운 변수나 모형은 같은 날짜의 MAE·MAPE와 발행 가능 범위를 함께 비교하며, 일부 거리의 악화와 미확정 결과도 공개합니다.', size=9.2, after=9)

doc.add_page_break()
section_banner(doc, source_banners[5])
add_text(doc, '현재 예측 실증은 수집·정합성을 확인한 가락·강서 경매/정산 이력을 사용합니다. 거래금액과 중량으로 가격을 집계하고 날짜별 관측 지원을 남깁니다. 공식 전국 목표와 순별 산식은 본선 자료에 맞춰 별도로 연결합니다.', size=10.2, after=6)
add_table(doc, [
    ['자료 번호·명칭', '사용할 정보 항목', '아이디어 활용 방법'],
    ['추가 제안(aT)\n가락·강서 경매/정산\n현재 실증 자료', '경락가격·거래량·품종·시장·거래일·정산 상태', '배추·무·양파·깐마늘의 가격 지연값·추세·계절성과 관측 지원을 구성해 10일 전망에 사용.'],
    ['추가 제안(aT)\n공개 전국 거래 파일\n목표 연결 참고', '시장·품목별 낙찰금액과 물량, 거래일, 시장 코드', '공식 전국 타깃과 시장·품종·단위·집계 정의를 대조. 현재 두 시장 실증과 구분.'],
    ['R01\n전국 경락가격·반입량\n본선 자료', '공식 가격·반입량·품종·단위·기간·집계 규칙', '공개 전국 계열과 행·산식·버전을 대조해 목표를 확정하고 다섯 품목·2순/3순으로 확장.'],
    ['R04·KOSIS\n조건부 추가자료', '기상 예보·관측, 생산·재배 통계와 발표시각', '원점 전에 공개된 값만 후보로 시험하고, 같은 목표 기준선보다 개선된 경우에만 채택.'],
], [1.55, 1.65, 2.39], size=8.7)
add_text(doc, 'KREI 자료는 제외합니다. 음수 정정 거래와 원문을 보존하고, 정상 빈 응답·API 실패·미완료 수집을 구분합니다. 양수 가격이 없다는 사실만으로 휴장이나 품목 무거래를 단정하지 않으며, 과거 수정본을 당시 공개판으로 가정하지 않습니다.', size=9.3, after=8)
add_label(doc, '변수 채택 원칙')
add_text(doc, '추가 변수·모형은 운영 E106과 보존 연구 대조군 E106-RHS를 함께 비교해 같은 원점·목표일·정답에서 평가합니다. MAE·MAPE, 원천 처리 상태와 발행 범위를 함께 보고 달력상 의존성과 다중비교를 반영해 개선·악화·미확정을 구분합니다. 실제 미래 비교는 거리별 공통 양수정답 30일 이상 확보 뒤 첫 월말을 추론 검토 시점으로 고정합니다. 이는 검정력 보장이 아니며, 그 전의 구현·후향 성과와 전향 검증을 구분합니다.', size=10.1, after=3)

doc.add_page_break()
section_banner(doc, source_banners[9])
issue_clock = datetime.fromisoformat(example_issue['issueAtKst']).strftime('%Y-%m-%d %H:%M KST')
target_dates = [datetime.fromisoformat(x['target']).strftime('%m/%d') for x in radish_predictions]
rounded_prices = [f"{x['predictedPrice']:,.1f}" for x in radish_predictions]
last_actual = f"{radish['lastObservedPrice']:,.1f}"
add_text(doc, f"가락·강서 E106 발행 예시 · 무 · {issue_clock} · 자료 최신일 {example_issue['sourceCutoffDate']}", size=10, bold=True, after=3)
add_text(doc, '품목 선택: 배추 · 무(선택) · 양파 · 깐마늘    건고추: 확장 계획', size=9.5, after=5)
add_table(doc, [
    ['선택 품목', '마지막 실제값', '최근 자료일', '발행시각'],
    ['무', f'{last_actual}원/kg', radish['lastObservedPriceDate'], issue_clock],
], [1.2, 1.45, 1.35, 1.65], size=8.6)
daily_rows = [
    ['예측가격\n원/kg'] + [f'D+{n}' for n in range(1, 11)],
    ['목표일\n(2026)'] + target_dates,
    ['무'] + rounded_prices,
]
add_table(doc, daily_rows, [0.82] + [0.594] * 10, size=7.8)
add_table(doc, [
    ['중기 범위', '화면 표시'],
    ['2순 후', '□ 원/kg · 공식 정의 대조 후 표시'],
    ['3순 후', '□ 원/kg · 공식 정의 대조 후 표시'],
], [1.25, 5.55], size=8.7)
add_text(doc, '표시 예측가는 0.1원/kg 단위 반올림입니다. 같은 시각·입력으로 기본 결합·E·E106을 비교 발행하고 그중 E106을 표시했습니다. 향후 정답이 아직 없어 미래 정확도는 미확정이며, 원천의 과거 최초 공개시각도 입증하지 못합니다. 농넷 본선의 공식 목표가격과 구분합니다.', size=9.1, after=5)
add_text(doc, '가격 추이:  실제 관측 ━━━     POC 예측 - - -     예측 범위는 검증을 마친 경우에만 선택 표시', size=9.2, after=5)
add_table(doc, [
    ['예측 근거와 상태', '표시할 내용'],
    ['현재 입력 정보', '최신 적격 가격일·관측 지원·원천 처리 상태'],
    ['후보 비교 결과', '기본 결합·E·E106의 동일쌍 오차와 검증 상태'],
], [2.2, 4.6], size=8.7)
add_text(doc, '화면 설명', size=10.5, bold=True, before=4, after=3)
add_text(doc, '이용자는 품목 선택 후 마지막 실제값과 D+1부터 D+10까지 날짜별 가격을 원/kg 단위로 비교하고, 예측 기준일·자료 최신일·발행시각을 확인합니다. 휴장·무거래·API 오류·자료 지연은 가격선에 섞지 않고 상태로 표시합니다. 품질 기준을 넘지 못하는 품목에는 새 숫자 전망을 내지 않습니다.', size=10, after=5)
add_text(doc, '2순·3순은 본선에서 정해지는 기간 경계와 공식 가격 집계에 맞춰 구현할 제안 기능입니다. 현재 화면의 해당 값은 비워 둡니다. 예측 배경은 검증된 자료와 모델 기여를 구분해 설명하고 원인으로 단정하지 않습니다. 사용자 의사결정 개선 효과는 향후 이용 평가로 확인해야 합니다.', size=10, after=3)

set_all_existing_fonts(doc)
doc.add_comment(
    rhs_update_p.runs,
    text='H/S 중단과 RHS 전달 진단 반영. 상세 근거: 연구보고서 §§12.9–12.10.',
    author='',
    initials='',
)
doc.save(OUTPUT)
scrub_package_metadata(OUTPUT)
print(OUTPUT)
print('template_sha256=' + hashlib.sha256(TEMPLATE.read_bytes()).hexdigest())
print('issue_at=' + example_issue['issueAtKst'])
print('issue_artifact_sha256=' + example_index['artifactSha256'])
